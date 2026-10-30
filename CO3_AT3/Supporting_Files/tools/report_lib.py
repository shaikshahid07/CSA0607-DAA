"""Small python-docx helper layer: institutional look (Times New Roman, white, plain tables)."""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SHOTS = os.path.join(ROOT, "Supporting_Files", "screenshots")
GRAPHS = os.path.join(ROOT, "Supporting_Files", "graphs")

STUDENT = "Shaik Shahid"
REG = "192424311"
COURSE = "CSA06 – Design and Analysis of Algorithms"
ASSESS = "CO3 – Case Study – Divide-and-Conquer Engine of a Retail Data Analytics Platform (KartIQ)"
FONT = "Times New Roman"


def _set_font(run, size=None, bold=None, italic=None, name=FONT, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def new_doc(header_text):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)      # A4
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.9)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.08
    for lvl, size in ((1, 14), (2, 12.5), (3, 11.5)):
        h = d.styles[f"Heading {lvl}"]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        h.element.rPr.rFonts.set(qn("w:cs"), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if h.element.rPr.rFonts.get(qn(a)) is not None:
                del h.element.rPr.rFonts.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    # header / footer
    hp = s.header.paragraphs[0]
    hp.text = ""
    _set_font(hp.add_run(header_text), 9.5, italic=True)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    _set_font(r, 9.5)
    _field(fp, "PAGE")
    return d


def _field(par, code):
    run = par.add_run()
    _set_font(run, 9.5)
    for t, txt in (("begin", None), (None, code), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def title_block(d, title, subtitle):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("Department of Computer Science and Engineering"), 12, bold=True)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    _set_font(p.add_run(title), 17, bold=True)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(subtitle), 11.5, italic=True)
    table(d, [["Student Name", STUDENT], ["Student Reg. No.", REG], ["Course Code & Name", COURSE],
              ["Assessment", ASSESS]], widths=[1.8, 4.4], header=False, bold_first_col=True, size=10.5)


def h(d, text, lvl=1):
    return d.add_heading(text, level=lvl)


def para(d, text, bold_prefix=None, italic=False, size=None, align=None, after=None):
    p = d.add_paragraph()
    if bold_prefix:
        _set_font(p.add_run(bold_prefix), size, bold=True)
    _set_font(p.add_run(text), size, italic=italic)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullets(d, items, size=None):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if isinstance(it, tuple):
            _set_font(p.add_run(it[0]), size, bold=True)
            _set_font(p.add_run(it[1]), size)
        else:
            _set_font(p.add_run(it), size)


def math(d, text):
    """Centered, italic line used for recurrences and formulas."""
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    _set_font(p.add_run(text), 11.5, italic=True)
    return p


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(d, rows, widths=None, header=True, bold_first_col=False, size=10, align_center_cols=(), keep=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            if widths:
                c.width = Inches(widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            is_h = header and i == 0
            _set_font(p.add_run(str(val)), size, bold=is_h or (bold_first_col and j == 0))
            if j in align_center_cols or is_h:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (j in align_center_cols or is_h) else None
            if is_h:
                _shade(c, "E7E6E6")
            elif bold_first_col and j == 0:
                _shade(c, "F2F2F2")
    for i, row in enumerate(t.rows):                      # rows never split; short tables stay on one page
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        if keep and len(rows) <= 12 and i < len(rows) - 1:
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    if header:   # repeat header row
        trPr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)
    sp = d.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    return t


def code(d, text, size=8.5):
    """Monospace block in a single shaded cell."""
    t = d.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    _shade(c, "F7F7F7")
    c.text = ""
    first = True
    for line in text.rstrip("\n").split("\n"):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        _set_font(p.add_run(line if line else " "), size, name="Courier New")
    sp = d.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)


def figure(d, path, caption, width=6.2):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width))
    c = d.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(c.add_run(caption), 10, italic=True)


def shot(d, name, caption, width=6.3):
    figure(d, os.path.join(SHOTS, name), caption, width)


def graph(d, name, caption, width=6.2):
    figure(d, os.path.join(GRAPHS, name), caption, width)


def read_text(rel):
    with open(os.path.join(ROOT, rel)) as f:
        return f.read()
